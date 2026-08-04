import io
import os

from app.core.constants import DocumentType


class FileService:
    SUPPORTED_FORMATS = {
        ".txt": DocumentType.TXT,
        ".md": DocumentType.MARKDOWN,
        ".docx": DocumentType.DOCX,
        ".pdf": DocumentType.PDF,
        ".rtf": DocumentType.RTF,
    }

    @staticmethod
    async def parse_file(filename: str, content: bytes) -> tuple[str, DocumentType]:
        ext = os.path.splitext(filename)[1].lower()
        doc_type = FileService.SUPPORTED_FORMATS.get(ext, DocumentType.TXT)

        if doc_type == DocumentType.TXT:
            return content.decode("utf-8", errors="replace"), doc_type
        elif doc_type == DocumentType.MARKDOWN:
            return content.decode("utf-8", errors="replace"), doc_type
        elif doc_type == DocumentType.DOCX:
            try:
                import docx
                doc = docx.Document(io.BytesIO(content))
                text = "\n".join(p.text for p in doc.paragraphs)
                return text, doc_type
            except ImportError:
                return content.decode("utf-8", errors="replace"), DocumentType.TXT
        elif doc_type == DocumentType.PDF:
            try:
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(content))
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
                return text, doc_type
            except ImportError:
                return content.decode("utf-8", errors="replace"), DocumentType.TXT
        elif doc_type == DocumentType.RTF:
            return content.decode("utf-8", errors="replace"), doc_type
        else:
            return content.decode("utf-8", errors="replace"), DocumentType.TXT

    @staticmethod
    async def export_file(content: str, format: DocumentType) -> tuple[bytes, str]:
        if format == DocumentType.TXT:
            return content.encode("utf-8"), "text/plain"
        elif format == DocumentType.MARKDOWN:
            return content.encode("utf-8"), "text/markdown"
        elif format == DocumentType.DOCX:
            try:
                import docx
                doc = docx.Document()
                for line in content.split("\n"):
                    if line.strip():
                        doc.add_paragraph(line)
                buf = io.BytesIO()
                doc.save(buf)
                return buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            except ImportError:
                return content.encode("utf-8"), "text/plain"
        elif format == DocumentType.PDF:
            try:
                from reportlab.lib.pagesizes import A4
                from reportlab.platypus import SimpleDocTemplate, Paragraph
                from reportlab.lib.styles import getSampleStyleSheet
                buf = io.BytesIO()
                doc = SimpleDocTemplate(buf, pagesize=A4)
                styles = getSampleStyleSheet()
                flowables = [Paragraph(line, styles["Normal"]) for line in content.split("\n") if line.strip()]
                doc.build(flowables)
                return buf.getvalue(), "application/pdf"
            except ImportError:
                return content.encode("utf-8"), "text/plain"
        else:
            return content.encode("utf-8"), "text/plain"
